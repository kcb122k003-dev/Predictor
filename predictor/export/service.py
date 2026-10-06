"""Export of predictions, formulations, evidence and backtest results (spec section 53).

CSV (one table), XLSX (all tables as sheets), PDF (readable report with charts) and JSON.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import Any

from sqlalchemy import select

from ..database.models import AnalysisRun, Course, PredictedQuestion
from ..prediction.ranking import DISCLAIMER
from ..services.context import AppContext
from ..services.results_service import ResultsService

TABLES = ("predictions", "questions", "models", "folds", "evidence", "excluded")


def _pct(v: Any) -> str:
    return "" if v is None else f"{100 * float(v):.0f}%"


class ExportService:
    def __init__(self, app: AppContext):
        self.app = app
        self.results = ResultsService(app)

    # ---------------------------------------------------------------- tables
    def tables(self, run_id: int) -> dict[str, list[dict[str, Any]]]:
        preds = self.results.predictions(run_id, "topic")
        models = self.results.models(run_id)
        with self.app.db.session() as s:
            forms = s.execute(select(PredictedQuestion).where(PredictedQuestion.run_id == run_id)
                              .order_by(PredictedQuestion.topic_id, PredictedQuestion.rank)).scalars().all()
            topic_label = {p["topic_id"]: p["label"] for p in preds}
            form_rows = [{"topic": topic_label.get(f.topic_id, ""), "label": "PREDICTED QUESTION FORMULATION",
                          "text": f.text, "format": f.question_type,
                          "marks": "" if f.marks_low is None else (f"{f.marks_low:g}" if f.marks_low == f.marks_high
                                                                   else f"{f.marks_low:g}-{f.marks_high:g}"),
                          "basis": f.basis, "note": (f.grounding or {}).get("note", ""),
                          "evidence_question_ids": " ".join(map(str, f.evidence_question_ids or []))} for f in forms]
        pred_rows = []
        evidence_rows = []
        for p in preds:
            facts = p["facts"] or {}
            tf = facts.get("type_forecast") or {}
            pred_rows.append({
                "rank": p["rank"], "topic": p["label"], "category": p["category"],
                "probability": _pct(p["probability"]) if p["calibrated"] else "",
                "probability_range": f"{_pct(p['prob_low'])}-{_pct(p['prob_high'])}" if p["calibrated"] else "",
                "relative_score": p["relative_score"], "confidence": p["confidence"],
                "recent_appearances": f"{facts.get('recent_appearances', '')}/{facts.get('recent_window', '')}",
                "total_appearances": f"{facts.get('appearances', '')}/{facts.get('exams', '')}",
                "last_appearance": facts.get("last_label") or "never",
                "typical_marks": "" if facts.get("marks_mean") is None else f"{facts['marks_min']:g}-{facts['marks_max']:g}",
                "likely_format": tf.get("format", ""), "syllabus_match": _pct(facts.get("mapping_confidence")),
            })
            for line in p["evidence"]:
                evidence_rows.append({"topic": p["label"], "kind": "evidence", "text": line})
            for line in p["why_not"]:
                evidence_rows.append({"topic": p["label"], "kind": "why not", "text": line})
            for group, value in (p["contributions"] or {}).items():
                evidence_rows.append({"topic": p["label"], "kind": "contribution", "text": f"{group}: {value:+.3f}"})
        model_rows = [{"layer": m["layer"], "model": m["display"], "enabled": m["enabled"], "selected": m["selected"],
                       "ndcg": m["metrics"].get("ndcg"), "ndcg_se": m["se"].get("ndcg"),
                       "recall": m["metrics"].get("recall"), "precision": m["metrics"].get("precision"),
                       "hit_rate": m["metrics"].get("hit_rate"), "mrr": m["metrics"].get("mrr"),
                       "hit@1": m["metrics"].get("hit@1"), "hit@3": m["metrics"].get("hit@3"),
                       "hit@5": m["metrics"].get("hit@5"), "hit@10": m["metrics"].get("hit@10"),
                       "why_disabled": m["gate_reason"]} for m in models["models"]]
        fold_rows = [{"model": f["model"], "target_exam": f["target"], "train_exams": f["train_exams"],
                      **{k: v for k, v in f["metrics"].items() if k in ("ndcg", "recall", "precision", "hit_rate", "mrr")}}
                     for f in models["folds"]]
        try:
            excluded = self.results.artifact(run_id, "excluded").get("groups", [])
        except KeyError:
            excluded = []
        excluded_rows = [{"question": g["label"], "years": ", ".join(g["years"]), "reason": g["reason"],
                          "nearest_topic": g.get("nearest_topic") or ""} for g in excluded]
        return {"predictions": pred_rows, "questions": form_rows, "models": model_rows, "folds": fold_rows,
                "evidence": evidence_rows, "excluded": excluded_rows}

    def csv(self, run_id: int, table: str = "predictions") -> bytes:
        if table not in TABLES:
            raise ValueError(f"table must be one of {', '.join(TABLES)}")
        rows = self.tables(run_id)[table]
        buf = io.StringIO()
        if rows:
            writer = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
            writer.writeheader()
            writer.writerows(rows)
        return buf.getvalue().encode("utf-8-sig")

    def xlsx(self, run_id: int) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        wb.remove(wb.active)
        info = wb.create_sheet("About")
        summary = self.results.run(run_id)["summary"]
        for row in self._about_rows(run_id, summary):
            info.append(row)
        info.column_dimensions["A"].width = 28
        info.column_dimensions["B"].width = 110
        for name, rows in self.tables(run_id).items():
            ws = wb.create_sheet(name.capitalize())
            if not rows:
                ws.append(["(no rows)"])
                continue
            headers = list(rows[0].keys())
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for r in rows:
                ws.append([r.get(h) for h in headers])
            for i, h in enumerate(headers, start=1):
                width = max(len(str(h)), *(len(str(r.get(h) or "")) for r in rows[:200]))
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(10, width + 2), 90)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def json(self, run_id: int) -> bytes:
        payload = {"run": self.results.run(run_id), "tables": self.tables(run_id)}
        for key in ("structure", "coverage", "ablation", "calibration", "type_forecast", "families", "papers",
                    "sufficiency", "syllabus_filter"):
            try:
                payload[key] = self.results.artifact(run_id, key)
            except KeyError:
                pass
        return json.dumps(payload, indent=2, default=str).encode("utf-8")

    def _about_rows(self, run_id: int, summary: dict[str, Any]) -> list[list[Any]]:
        with self.app.db.session() as s:
            run = s.get(AnalysisRun, run_id)
            course = s.get(Course, run.course_id)
            name = course.name
        return [
            ["Course", name], ["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
            ["Important", DISCLAIMER], ["Exams analysed", summary.get("exams")],
            ["Questions", summary.get("questions")], ["Topics", summary.get("topics")],
            ["Model used", summary.get("selected_display")], ["Why this model", summary.get("selection_reason")],
            ["Probabilities", summary.get("calibration_reason")],
            ["Notes", " ".join(summary.get("notes") or [])],
        ]

    # -------------------------------------------------------------------- PDF
    def pdf(self, run_id: int) -> bytes:
        from reportlab.graphics.charts.barcharts import HorizontalBarChart
        from reportlab.graphics.shapes import Drawing
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)
        from xml.sax.saxutils import escape

        summary = self.results.run(run_id)["summary"]
        tables = self.tables(run_id)
        preds = self.results.predictions(run_id, "topic")
        size = letter if str(self.app.settings.export.pdf_page_size).lower() == "letter" else A4
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=size, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=16 * mm,
                                bottomMargin=16 * mm, title="Exam prediction report")
        styles = getSampleStyleSheet()
        small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=8.5, leading=11)
        body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=12.5)
        warn = ParagraphStyle("warn", parent=body, textColor=colors.HexColor("#8a4b00"))
        story: list[Any] = []
        about = dict((r[0], r[1]) for r in self._about_rows(run_id, summary))

        def p(text: str, style=body):
            return Paragraph(escape(str(text)), style)

        story += [Paragraph("Exam prediction report", styles["Title"]),
                  p(f"Course: {about['Course']}   |   Generated {about['Generated']}"), Spacer(1, 4),
                  p(DISCLAIMER, warn), Spacer(1, 6)]
        story.append(Paragraph("Summary", styles["Heading2"]))
        for label in ("Exams analysed", "Questions", "Topics", "Model used", "Why this model", "Probabilities"):
            story.append(p(f"{label}: {about[label]}"))
        for note in summary.get("notes") or []:
            story.append(p(note, warn))
        suff = (summary.get("sufficiency") or {}).get("message")
        if suff:
            story.append(p(suff, warn))

        story.append(Paragraph("Ranked topics", styles["Heading2"]))
        calibrated = summary.get("calibrated")
        head = ["#", "Topic", "Priority", "Probability" if calibrated else "Relative score", "Confidence",
                "Recent", "Last seen", "Format"]
        data = [head]
        for r in tables["predictions"][:40]:
            value = f"{r['probability']} ({r['probability_range']})" if calibrated else f"{r['relative_score']:.2f}"
            data.append([r["rank"], Paragraph(escape(r["topic"]), small), r["category"].replace(" Priority", ""),
                         value, r["confidence"], r["recent_appearances"], r["last_appearance"], r["likely_format"]])
        t = Table(data, repeatRows=1, colWidths=[8 * mm, 62 * mm, 22 * mm, 26 * mm, 18 * mm, 14 * mm, 24 * mm, 18 * mm])
        t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.5), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8ecf2")),
                               ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c8ced8")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)

        top = [x for x in preds if x["rank"] <= 12]
        if top:
            story.append(Spacer(1, 8))
            drawing = Drawing(170 * mm, 6 * mm * len(top) + 20)
            chart = HorizontalBarChart()
            chart.x, chart.y = 60 * mm, 10
            chart.width, chart.height = 100 * mm, 6 * mm * len(top)
            vals = [x["probability"] if calibrated and x["probability"] is not None else x["relative_score"] or 0
                    for x in reversed(top)]
            chart.data = [vals]
            chart.categoryAxis.categoryNames = [x["label"][:42] for x in reversed(top)]
            chart.categoryAxis.labels.fontSize = 6.5
            chart.valueAxis.valueMin, chart.valueAxis.valueMax = 0, 1
            chart.valueAxis.labels.fontSize = 7
            chart.bars[0].fillColor = colors.HexColor("#3b6ea8")
            drawing.add(chart)
            story.append(drawing)

        story.append(PageBreak())
        story.append(Paragraph("Why the top topics ranked highly", styles["Heading2"]))
        for x in preds[:10]:
            story.append(Paragraph(escape(f"{x['rank']}. {x['label']} - {x['category']}"), styles["Heading4"]))
            for line in x["evidence"]:
                story.append(p(f"- {line}", small))
            forms = [f for f in tables["questions"] if f["topic"] == x["label"]]
            for f in forms[:3]:
                story.append(p(f"PREDICTED QUESTION FORMULATION ({f['format']}, marks {f['marks'] or '?'}): {f['text']}", small))
                if f["note"]:
                    story.append(p(f"  {f['note']}", small))
        low = [x for x in preds if x["why_not"]][:8]
        if low:
            story.append(Paragraph("Why some topics ranked low", styles["Heading2"]))
            for x in low:
                story.append(p(f"{x['label']}: {' '.join(x['why_not'])}", small))

        story.append(PageBreak())
        story.append(Paragraph("Model comparison (time-ordered backtest)", styles["Heading2"]))
        mrows = [["Model", "Selected", "NDCG", "Recall", "Hit@3", "Hit@5", "Status"]]
        for m in tables["models"]:
            if m["layer"] != "topic":
                continue
            fmt = lambda v: "" if v is None else f"{v:.3f}"
            mrows.append([Paragraph(escape(m["model"]), small), "yes" if m["selected"] else "", fmt(m["ndcg"]),
                          fmt(m["recall"]), fmt(m["hit@3"]), fmt(m["hit@5"]),
                          Paragraph(escape("enabled" if m["enabled"] else m["why_disabled"]), small)])
        mt = Table(mrows, repeatRows=1, colWidths=[46 * mm, 15 * mm, 15 * mm, 15 * mm, 15 * mm, 15 * mm, 60 * mm])
        mt.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 7.5), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8ecf2")),
                                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c8ced8")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(mt)
        try:
            ablation = self.results.artifact(run_id, "ablation")
        except KeyError:
            ablation = {}
        story.append(Paragraph("Ablation study", styles["Heading3"]))
        if ablation.get("available"):
            for row in ablation.get("staged", []):
                delta = "" if row.get("delta") is None else f" (change {row['delta']:+.3f} +/- {row.get('delta_se') or 0:.3f})"
                story.append(p(f"{row['variant']}: {row['mean']}{delta}", small))
        else:
            story.append(p(ablation.get("reason", "Not available."), small))
        if tables["excluded"]:
            story.append(Paragraph("Excluded: outside the current syllabus", styles["Heading2"]))
            for e in tables["excluded"]:
                story.append(p(f"{e['question']} ({e['years']}). {e['reason']}", small))
        doc.build(story)
        return buf.getvalue()
