"""Ingestion of real file formats: digital PDF, scanned PDF, DOCX (with Word numbering), images, text."""

import io

import pytest
from PIL import Image, ImageDraw, ImageFont

from predictor.ingestion.extract import UnsupportedFileError, extract_document, sniff_type
from predictor.ocr.engine import OcrEngine
from predictor.parsing.exam_parser import ExamParser
from predictor.services.context import AppContext
from predictor.services.course_service import CourseService
from predictor.services.ingest_service import IngestService

from tests.conftest import requires_tesseract

PAPER_LINES = [
    "Final Examination 2021 Spring", "Subject: Thermodynamics   Full Marks: 40",
    "1. a) Define entropy and state its units. [4]",
    "   b) Derive the Clausius inequality for a cyclic process. [6]",
    "2. Calculate the efficiency of a Carnot engine working between 600 K and 300 K. [5]",
]


def _font(size=34):
    for path in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "C:/Windows/Fonts/arial.ttf"):
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def _paper_image(rotate: float = 0.0) -> Image.Image:
    img = Image.new("RGB", (2000, 800), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(PAPER_LINES):
        draw.text((60, 60 + 130 * i), line, fill="black", font=_font())
    return img.rotate(rotate, expand=True, fillcolor="white") if rotate else img


def test_digital_pdf(tmp_path, settings):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    path = tmp_path / "paper.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    for i, line in enumerate(PAPER_LINES):
        c.drawString(50, 800 - 18 * i, line)
    c.save()
    res = extract_document(path, settings, None)
    assert res.pages[0].method == "native"
    parsed = ExamParser(settings).parse(res.pages, path.name)
    labels = [n.path_label for n in parsed.all_nodes()]
    assert labels == ["1", "1(a)", "1(b)", "2"]
    assert parsed.metadata.get("year") == 2021


def test_docx_with_word_numbering(tmp_path, settings):
    import docx

    d = docx.Document()
    d.add_paragraph("Examination 2020")
    d.add_paragraph("Explain the first law of thermodynamics. [5]", style="List Number")
    d.add_paragraph("Derive the Maxwell relations. [8]", style="List Number")
    table = d.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "3."
    table.rows[0].cells[1].text = "Calculate the COP of a refrigerator working between 250 K and 300 K."
    table.rows[0].cells[2].text = "[6]"
    path = tmp_path / "paper.docx"
    d.save(path)
    res = extract_document(path, settings, None)
    text = res.pages[0].text
    assert "1. Explain the first law" in text and "2. Derive the Maxwell" in text  # numbers come from numbering.xml
    parsed = ExamParser(settings).parse(res.pages, path.name)
    assert [q.label for q in parsed.questions] == ["1", "2", "3"]
    assert [q.marks for q in parsed.questions] == [5, 8, 6]


@requires_tesseract
def test_scanned_pdf_uses_ocr(tmp_path, settings):
    import pymupdf

    buf = io.BytesIO()
    _paper_image(rotate=1.5).save(buf, format="PNG")
    path = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(pymupdf.Rect(20, 20, 575, 260), stream=buf.getvalue())
    doc.save(str(path))
    res = extract_document(path, settings, OcrEngine(settings))
    assert res.pages[0].method == "ocr" and res.pages[0].confidence > 60
    parsed = ExamParser(settings).parse(res.pages, path.name)
    assert [q.label for q in parsed.questions] == ["1", "2"]
    assert parsed.questions[0].children[1].marks == 6


@requires_tesseract
def test_image_ocr_with_quality_flags(tmp_path, settings):
    path = tmp_path / "photo.png"
    _paper_image().save(path)
    res = extract_document(path, settings, OcrEngine(settings))
    assert "Carnot" in res.pages[0].text
    assert res.pages[0].confidence is not None


def test_image_without_ocr_is_reported(tmp_path, settings):
    path = tmp_path / "photo.png"
    _paper_image().save(path)
    off = settings.merged({"ocr": {"enabled": False}})
    with pytest.raises(RuntimeError, match="OCR"):
        extract_document(path, off, OcrEngine(off))


def test_unsupported_and_duplicate_files(tmp_path):
    old = tmp_path / "old.doc"
    old.write_bytes(b"\xd0\xcf\x11\xe0legacy")
    with pytest.raises(UnsupportedFileError):
        sniff_type(old)
    app = AppContext.create(tmp_path / "data", log=False)
    cid = CourseService(app).create("C")
    paper = tmp_path / "p.txt"
    paper.write_text("\n".join(PAPER_LINES), encoding="utf-8")
    service = IngestService(app)
    first = service.add_file(cid, paper, "p.txt", "exam")
    second = service.add_file(cid, paper, "copy.txt", "exam")
    assert not first.duplicate and second.duplicate and second.file_id == first.file_id
    # The same file in a different course is not a duplicate there (courses are independent).
    other = CourseService(app).create("D")
    assert not service.add_file(other, paper, "p.txt", "exam").duplicate
